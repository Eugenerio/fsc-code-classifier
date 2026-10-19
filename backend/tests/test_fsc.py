from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.fsc import GROUP_TITLES, FscCatalog, load_catalog


def write_xlsx(path: Path, rows: list[tuple], header: tuple = ("FSC", "Description", "SOS RIC", "Marker")) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FSC"
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


@pytest.fixture(scope="module")
def catalog() -> FscCatalog:
    return load_catalog()


def test_real_catalog_shape(catalog: FscCatalog) -> None:
    assert len(catalog.codes) == 580
    assert [c.code for c in catalog.codes] == sorted(c.code for c in catalog.codes)
    assert all(len(c.code) == 4 and c.code.isdigit() for c in catalog.codes)
    assert all(c.description.strip() and c.sos_ric for c in catalog.codes)
    assert set(catalog.by_code) == {c.code for c in catalog.codes}


def test_lookup(catalog: FscCatalog) -> None:
    screws = catalog.get("5305")
    assert screws is not None
    assert (screws.description, screws.group, screws.group_name) == ("Screws", "53", "Hardware and Abrasives")
    assert catalog.get("0000") is None
    assert catalog.get("1005").description == "Guns, thru 30 mm"  # numeric cell 1005 kept its 4 digits
    assert catalog.get("5345").marker == "**"
    assert catalog.get("5305").marker is None


def test_every_group_prefix_has_an_official_title(catalog: FscCatalog) -> None:
    prefixes = {c.code[:2] for c in catalog.codes}
    assert prefixes <= GROUP_TITLES.keys(), f"missing group titles: {sorted(prefixes - GROUP_TITLES.keys())}"
    assert catalog.groups == {g: GROUP_TITLES[g] for g in sorted(prefixes)}
    assert catalog.groups["15"] == "Aircraft and Airframe Structural Components"
    assert catalog.groups["99"] == "Miscellaneous"


def test_in_group(catalog: FscCatalog) -> None:
    hardware = catalog.in_group("53")
    assert hardware and {c.code[:2] for c in hardware} == {"53"}
    assert "5305" in {c.code for c in hardware}
    assert sum(len(catalog.in_group(g)) for g in catalog.groups) == len(catalog.codes)
    assert catalog.in_group("00") == []


def test_prompt_lines_are_grouped_and_compact(catalog: FscCatalog) -> None:
    lines = catalog.prompt_lines().splitlines()
    headers = [ln for ln in lines if ln.startswith("## ")]
    assert len(headers) == len(catalog.groups)
    assert "## 53 Hardware and Abrasives" in headers
    start = lines.index("## 53 Hardware and Abrasives")
    assert lines[start + 1] == "5305 | Screws | General & Industrial"
    assert lines[start + 2] == "5306 | Bolts | General & Industrial"
    code_lines = [ln for ln in lines if not ln.startswith("## ")]
    assert len(code_lines) == 580
    assert all(ln[4:7] == " | " for ln in code_lines)
    assert "7510 | Office Supplies | GSA Office & Paper | commercial" in code_lines
    assert "8905 | Meat, Poultry and Fish | Subsistence | subsistence" in code_lines
    assert len(catalog.prompt_lines()) < 45_000  # ~10k tokens: must fit comfortably in one prompt


def test_load_catalog_is_cached() -> None:
    assert load_catalog() is load_catalog()


def test_loader_normalises_codes_and_cells(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "fsc.xlsx",
        [
            (5305, "  Screws  ", "S9I/KZ", None),
            (" 5306 ", "Bolts\nand   studs", " S9I/KZ ", "**"),
            (4710.0, "Pipe and Tube", "S9G/CX", None),
            (805, "Short code", "S9G/CX", None),
            (None, None, None, None),  # blank row
            ("  ", None, None, None),  # whitespace-only code
        ],
    )
    cat = load_catalog(path)
    assert [c.code for c in cat.codes] == ["0805", "4710", "5305", "5306"]
    assert cat.get("5305").description == "Screws"
    assert cat.get("5306").description == "Bolts and studs"
    assert cat.get("5306").sos_ric == "S9I/KZ"
    assert cat.get("5306").marker == "**"
    assert cat.get("0805").group_name == "Group 08"  # unknown prefix falls back instead of crashing


def test_loader_reads_columns_by_header_name(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "reordered.xlsx",
        [("Screws", "S9I/KZ", None, 5305)],
        header=("Description", "SOS RIC", "Marker", "FSC"),
    )
    assert load_catalog(path).get("5305").description == "Screws"


def test_loader_rejects_duplicates(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "dupes.xlsx",
        [(5305, "Screws", "S9I/KZ", None), (5306, "Bolts", "S9I/KZ", None), ("5305", "Screws again", "S9I/KZ", None)],
    )
    with pytest.raises(ValueError, match=r"duplicate FSC codes.*5305"):
        load_catalog(path)


@pytest.mark.parametrize("bad", ["53A5", 123456, "abc"])
def test_loader_rejects_malformed_codes(tmp_path: Path, bad: object) -> None:
    path = write_xlsx(tmp_path / "bad.xlsx", [(bad, "Whatever", "S9I/KZ", None)])
    with pytest.raises(ValueError, match="invalid FSC code"):
        load_catalog(path)


# --------------------------------------------------------------------------- Legend: supply domain, agency, marker


def test_legend_gives_every_class_an_agency_and_supply_domain(catalog: FscCatalog) -> None:
    assert len(catalog.domains) == 14
    assert sum(len(d.codes) for d in catalog.domains.values()) == 580
    assert {d.agency for d in catalog.domains.values()} == {"DLA", "GSA"}
    assert all(c.agency in ("DLA", "GSA") and c.supply_domain for c in catalog.codes)
    screws = catalog.get("5305")
    assert (screws.sos_ric, screws.agency, screws.supply_domain) == ("S9I/KZ", "DLA", "General & Industrial")
    paper = catalog.get("7510")
    assert (paper.agency, paper.supply_domain) == ("GSA", "GSA Office & Paper")
    columbus = catalog.domains["S9E/TX"]
    assert (columbus.agency, columbus.label, len(columbus.codes)) == ("DLA", "Electronics", 68)
    assert "Defense Supply Center, Columbus" in columbus.supplier
    assert catalog.domain_of("5905") is columbus
    assert catalog.domain_of("0000") is None
    assert len({d.label for d in catalog.domains.values()}) == 14  # labels double as Jev choice options: unique


def test_marker_counts_notes_and_kind(catalog: FscCatalog) -> None:
    kinds = [c.marker_kind for c in catalog.codes]
    assert kinds.count("commercial") == 39 and kinds.count("subsistence") == 16
    assert all((c.marker_note is not None) == (c.marker in ("*", "**")) for c in catalog.codes)
    assert catalog.get("7510").marker_note == "supported by both GSA and DLA (commercial goods)"
    assert catalog.get("8905").marker_note == "subsistence / live-animal class"
    assert catalog.get("5305").marker_note is None and catalog.get("5305").marker_kind is None


def test_catalog_text_is_the_enriched_embedding_text(catalog: FscCatalog) -> None:
    from app.fsc import catalog_text

    assert catalog_text(catalog.get("5305")) == (
        "Screws. Group: Hardware and Abrasives. Supplied by: General & Industrial (DLA)."
    )
    assert catalog_text(catalog.get("7510")).endswith("Supported by both GSA and DLA (commercial goods).")


def test_catalog_without_legend_falls_back_to_static_labels(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "nolegend.xlsx", [(5305, "Screws", "S9I/KZ", None), (7510, "Office", "2FY/75", "**")])
    cat = load_catalog(path)
    assert (cat.get("5305").agency, cat.get("5305").supply_domain) == ("DLA", "General & Industrial")
    assert (cat.get("7510").agency, cat.get("7510").supply_domain) == ("GSA", "GSA Office & Paper")
    assert set(cat.domains) == {"S9I/KZ", "2FY/75"}


def test_unknown_ric_uses_legend_parenthesised_text(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FSC"
    ws.append(("FSC", "Description", "SOS RIC", "Marker"))
    ws.append((5305, "Screws", "ZZZ", None))
    legend = wb.create_sheet("Legend")
    legend.append(("SOS RIC", "Agency", "Source of supply (as named in the PDF legend)", "FSC count"))
    legend.append(("ZZZ", "DLA", "Defense Supply Center, Mars (Rocketry, formerly Space Supply)", 1))
    legend.append((None, None, "Total", 1))
    wb.save(tmp_path / "custom.xlsx")
    cat = load_catalog(tmp_path / "custom.xlsx")
    assert cat.get("5305").supply_domain == "Rocketry" and cat.get("5305").agency == "DLA"
