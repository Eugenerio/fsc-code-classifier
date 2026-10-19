"""FSC catalog: loads the DLA/GSA assignment sheet and exposes it as typed, indexed objects."""

from __future__ import annotations

from collections import Counter, defaultdict
from functools import cached_property, lru_cache
from pathlib import Path

import openpyxl
from pydantic import BaseModel

from app.config import get_settings
from app.schemas import FscCode

# Official Federal Supply Classification group titles (2-digit prefix -> title).
GROUP_TITLES: dict[str, str] = {
    "10": "Weapons",
    "11": "Nuclear Ordnance",
    "12": "Fire Control Equipment",
    "13": "Ammunition and Explosives",
    "14": "Guided Missiles",
    "15": "Aircraft and Airframe Structural Components",
    "16": "Aircraft Components and Accessories",
    "17": "Aircraft Launching, Landing, and Ground Handling Equipment",
    "18": "Space Vehicles",
    "19": "Ships, Small Craft, Pontoons, and Floating Docks",
    "20": "Ship and Marine Equipment",
    "22": "Railway Equipment",
    "23": "Ground Effect Vehicles, Motor Vehicles, Trailers, and Cycles",
    "24": "Tractors",
    "25": "Vehicular Equipment Components",
    "26": "Tires and Tubes",
    "28": "Engines, Turbines, and Components",
    "29": "Engine Accessories",
    "30": "Mechanical Power Transmission Equipment",
    "31": "Bearings",
    "32": "Woodworking Machinery and Equipment",
    "34": "Metalworking Machinery",
    "35": "Service and Trade Equipment",
    "36": "Special Industry Machinery",
    "37": "Agricultural Machinery and Equipment",
    "38": "Construction, Mining, Excavating, and Highway Maintenance Equipment",
    "39": "Materials Handling Equipment",
    "40": "Rope, Cable, Chain, and Fittings",
    "41": "Refrigeration, Air Conditioning, and Air Circulating Equipment",
    "42": "Fire Fighting, Rescue, and Safety Equipment; and Environmental Protection Equipment and Materials",
    "43": "Pumps and Compressors",
    "44": "Furnace, Steam Plant, and Drying Equipment; and Nuclear Reactors",
    "45": "Plumbing, Heating, and Waste Disposal Equipment",
    "46": "Water Purification and Sewage Treatment Equipment",
    "47": "Pipe, Tubing, Hose, and Fittings",
    "48": "Valves",
    "49": "Maintenance and Repair Shop Equipment",
    "51": "Hand Tools",
    "52": "Measuring Tools",
    "53": "Hardware and Abrasives",
    "54": "Prefabricated Structures and Scaffolding",
    "55": "Lumber, Millwork, Plywood, and Veneer",
    "56": "Construction and Building Materials",
    "58": "Communication, Detection, and Coherent Radiation Equipment",
    "59": "Electrical and Electronic Equipment Components",
    "60": "Fiber Optics Materials, Components, Assemblies, and Accessories",
    "61": "Electric Wire, and Power and Distribution Equipment",
    "62": "Lighting Fixtures and Lamps",
    "63": "Alarm, Signal, and Security Detection Systems",
    "65": "Medical, Dental, and Veterinary Equipment and Supplies",
    "66": "Instruments and Laboratory Equipment",
    "67": "Photographic Equipment",
    "68": "Chemicals and Chemical Products",
    "69": "Training Aids and Devices",
    "70": "General Purpose Information Technology Equipment",
    "71": "Furniture",
    "72": "Household and Commercial Furnishings and Appliances",
    "73": "Food Preparation and Serving Equipment",
    "74": "Office Machines, Text Processing Systems and Visible Record Equipment",
    "75": "Office Supplies and Devices",
    "76": "Books, Maps, and Other Publications",
    "77": "Musical Instruments, Phonographs, and Home-Type Radios",
    "78": "Recreational and Athletic Equipment",
    "79": "Cleaning Equipment and Supplies",
    "80": "Brushes, Paints, Sealers, and Adhesives",
    "81": "Containers, Packaging, and Packing Supplies",
    "83": "Textiles, Leather, Furs, Apparel and Shoe Findings, Tents and Flags",
    "84": "Clothing, Individual Equipment, and Insignia",
    "85": "Toiletries",
    "87": "Agricultural Supplies",
    "88": "Live Animals",
    "89": "Subsistence",
    "91": "Fuels, Lubricants, Oils, and Waxes",
    "93": "Nonmetallic Fabricated Materials",
    "94": "Nonmetallic Crude Materials",
    "95": "Metal Bars, Sheets, and Shapes",
    "96": "Ores, Minerals, and Their Primary Products",
    "99": "Miscellaneous",
}


# SOS RIC (source of supply) -> short supply-domain label. The Legend sheet only has long supplier names, so this
# static map gives clean labels; an unknown RIC falls back to the Legend's parenthesised text or the supplier name.
DOMAIN_LABELS: dict[str, str] = {
    "S9C/AX": "Construction",
    "S9E/TX": "Electronics",
    "S9G/CX": "Aviation & General",
    "S9I/KZ": "General & Industrial",
    "S9M/KX": "Medical",
    "S9T/CY": "Clothing & Textiles",
    "S9S/S9P": "Subsistence",
    "S9F/KY": "Fuel",
    "HM8": "Mapping",
    "2FY/75": "GSA Office & Paper",
    "3FN/75": "GSA Furniture",
    "6FE/75": "GSA Tools & Appliances",
    "7FX/75": "GSA General Products",
    "FAA/75": "GSA Vehicles",
}
MARKER_NOTES: dict[str, str] = {
    "**": "supported by both GSA and DLA (commercial goods)",
    "*": "subsistence / live-animal class",
}


class DomainInfo(BaseModel):
    """One SOS RIC from the Legend sheet: the buying centre that owns a set of classes."""

    ric: str
    agency: str  # "DLA" | "GSA"
    supplier: str  # Legend text, e.g. "Defense Supply Center, Columbus (Construction, ...)"
    label: str  # short, e.g. "Construction"
    codes: list[str] = []


def _agency_of(ric: str) -> str:
    return "GSA" if ric.endswith("/75") else "DLA"


def _fallback_label(ric: str, supplier: str) -> str:
    if "(" in supplier and ")" in supplier:
        inner = supplier[supplier.index("(") + 1 : supplier.rindex(")")]
        return inner.split(",")[0].strip() or ric
    return supplier or ric


def catalog_text(fsc: FscCode) -> str:
    """Enriched, human-readable text of one class: what the embedding channel indexes."""
    parts = [f"{fsc.description}.", f"Group: {fsc.group_name}."]
    if fsc.supply_domain:
        parts.append(f"Supplied by: {fsc.supply_domain}" + (f" ({fsc.agency})." if fsc.agency else "."))
    if fsc.marker_note:
        parts.append(f"{fsc.marker_note[0].upper()}{fsc.marker_note[1:]}.")
    return " ".join(parts)


class FscCatalog:
    def __init__(self, codes: list[FscCode], legend: dict[str, tuple[str, str]] | None = None) -> None:
        """`legend` maps SOS RIC -> (agency, supplier text); without it agency and label come from static tables."""
        legend = legend or {}
        self.codes = sorted((self._enrich(c, legend) for c in codes), key=lambda c: c.code)
        self.by_code = {c.code: c for c in self.codes}
        self._by_group: dict[str, list[FscCode]] = defaultdict(list)
        for c in self.codes:
            self._by_group[c.group].append(c)
        self.groups = {g: items[0].group_name for g, items in self._by_group.items()}
        self.domains: dict[str, DomainInfo] = {}
        for c in self.codes:
            info = self.domains.get(c.sos_ric)
            if info is None:
                supplier = legend.get(c.sos_ric, ("", ""))[1]
                info = self.domains[c.sos_ric] = DomainInfo(
                    ric=c.sos_ric, agency=c.agency, supplier=supplier, label=c.supply_domain, codes=[]
                )
            info.codes.append(c.code)

    @staticmethod
    def _enrich(c: FscCode, legend: dict[str, tuple[str, str]]) -> FscCode:
        agency, supplier = legend.get(c.sos_ric, ("", ""))
        label = DOMAIN_LABELS.get(c.sos_ric) or _fallback_label(c.sos_ric, supplier)
        return c.model_copy(
            update={
                "agency": c.agency or agency or _agency_of(c.sos_ric),
                "supply_domain": c.supply_domain or label,
                "marker_note": c.marker_note or MARKER_NOTES.get(c.marker or ""),
            }
        )

    def domain_of(self, code: str) -> DomainInfo | None:
        fsc = self.by_code.get(code)
        return self.domains.get(fsc.sos_ric) if fsc else None

    def get(self, code: str) -> FscCode | None:
        return self.by_code.get(code)

    def in_group(self, group: str) -> list[FscCode]:
        return list(self._by_group.get(group, []))

    @cached_property
    def _prompt_text(self) -> str:
        lines: list[str] = []
        for group, title in self.groups.items():
            lines.append(f"## {group} {title}")
            lines.extend(self._prompt_line(c) for c in self._by_group[group])
        return "\n".join(lines)

    @staticmethod
    def _prompt_line(c: FscCode) -> str:
        cells = [c.code, c.description, c.supply_domain]
        if c.marker_kind:
            cells.append(c.marker_kind)
        return " | ".join(cell for cell in cells if cell)

    def prompt_lines(self) -> str:
        """Compact listing for LLM prompts: `7510 | Office Supplies | GSA Office & Paper | commercial` per code,
        grouped under `## 75 Office Supplies and Devices` headers (the marker column only for marked classes)."""
        return self._prompt_text


def _normalize_code(raw: object) -> str:
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)  # spreadsheets often hand back 5305.0
    code = str(raw).strip().zfill(4)
    if len(code) != 4 or not code.isdigit():
        raise ValueError(f"invalid FSC code {raw!r}")
    return code


def _read_rows(path: Path) -> list[FscCode]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb["FSC"].iter_rows(values_only=True)
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
        col = {name: header.index(name) for name in ("fsc", "description", "sos ric", "marker")}
        codes: list[FscCode] = []
        for row in rows:
            raw_code = row[col["fsc"]]
            if raw_code is None or not str(raw_code).strip():
                continue
            code = _normalize_code(raw_code)
            group = code[:2]
            marker = row[col["marker"]]
            codes.append(
                FscCode(
                    code=code,
                    description=" ".join(str(row[col["description"]] or "").split()),
                    sos_ric=str(row[col["sos ric"]] or "").strip(),
                    group=group,
                    group_name=GROUP_TITLES.get(group, f"Group {group}"),
                    marker=str(marker).strip() if marker else None,
                )
            )
        return codes
    finally:
        wb.close()


def _read_legend(path: Path) -> dict[str, tuple[str, str]]:
    """SOS RIC -> (agency, supplier text) from the Legend sheet; {} when the sheet is missing or unreadable."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Legend" not in wb.sheetnames:
            return {}
        rows = wb["Legend"].iter_rows(values_only=True)
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows, ())]
        try:
            ric_col = header.index("sos ric")
            agency_col = header.index("agency")
            supplier_col = next(i for i, h in enumerate(header) if h.startswith("source of supply"))
        except (ValueError, StopIteration):
            return {}
        legend: dict[str, tuple[str, str]] = {}
        for row in rows:
            ric = str(row[ric_col] or "").strip()
            if not ric:  # the table ends at the "Total" row, whose RIC cell is blank
                break
            legend[ric] = (str(row[agency_col] or "").strip(), " ".join(str(row[supplier_col] or "").split()))
        return legend
    finally:
        wb.close()


@lru_cache
def load_catalog(path: Path | None = None) -> FscCatalog:
    path = path or get_settings().fsc_xlsx_path
    codes = _read_rows(path)
    dupes = sorted(code for code, n in Counter(c.code for c in codes).items() if n > 1)
    if dupes:
        raise ValueError(f"duplicate FSC codes in catalog: {', '.join(dupes)}")
    return FscCatalog(codes, _read_legend(path))
